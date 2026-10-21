"""Build a PowerPoint presentation from rendered or extracted PDF content."""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Tuple

import fitz
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from app.core.object_extractor import (
    ExtractedDrawing,
    ExtractedImage,
    ExtractedTextBlock,
    extract_page,
)
from app.core.pdf_renderer import render_page
from app.models.conversion_options import ConversionOptions, SlideRatio

EMU_PER_POINT = 12700
DEFAULT_SLIDE_WIDTH_IN = 13.3333
DEFAULT_SLIDE_HEIGHT_IN = 7.5
MIN_SLIDE_IN = 1.0
MIN_FONT_PT = 1.0
MAX_FONT_PT = 400.0


@dataclass
class BuildSummary:
    text_shapes: int = 0
    images: int = 0
    shapes: int = 0
    warnings: List[str] = field(default_factory=list)


@dataclass
class SlideMapping:
    """Maps PDF points to EMU on a slide with letterboxing."""

    slide_width_emu: int
    slide_height_emu: int
    page_width: float
    page_height: float

    @property
    def scale(self) -> float:
        """EMU per PDF point (isotropic, content is never stretched)."""
        return min(
            self.slide_width_emu / max(self.page_width, 0.01),
            self.slide_height_emu / max(self.page_height, 0.01),
        )

    @property
    def offset_x(self) -> float:
        return (self.slide_width_emu - self.page_width * self.scale) / 2.0

    @property
    def offset_y(self) -> float:
        return (self.slide_height_emu - self.page_height * self.scale) / 2.0

    def x(self, value: float) -> int:
        return int(round(self.offset_x + value * self.scale))

    def y(self, value: float) -> int:
        return int(round(self.offset_y + value * self.scale))

    def rect(self, bbox: Tuple[float, float, float, float]) -> Tuple[int, int, int, int]:
        left = self.x(bbox[0])
        top = self.y(bbox[1])
        width = max(1, int(round((bbox[2] - bbox[0]) * self.scale)))
        height = max(1, int(round((bbox[3] - bbox[1]) * self.scale)))
        return left, top, width, height

    def font_pt(self, pdf_size: float) -> float:
        return max(MIN_FONT_PT, min(MAX_FONT_PT, pdf_size * self.scale / EMU_PER_POINT))


def compute_slide_size(
    page_width: float, page_height: float, ratio: SlideRatio
) -> Tuple[int, int]:
    """Return slide width/height in EMU for the requested ratio."""

    if ratio is SlideRatio.WIDESCREEN:
        return Inches(DEFAULT_SLIDE_WIDTH_IN), Inches(DEFAULT_SLIDE_HEIGHT_IN)
    if ratio is SlideRatio.STANDARD:
        return Inches(10.0), Inches(DEFAULT_SLIDE_HEIGHT_IN)

    aspect = max(0.05, min(20.0, page_width / max(page_height, 0.01)))
    if aspect >= DEFAULT_SLIDE_WIDTH_IN / DEFAULT_SLIDE_HEIGHT_IN:
        width_in = DEFAULT_SLIDE_WIDTH_IN
        height_in = width_in / aspect
    else:
        height_in = DEFAULT_SLIDE_HEIGHT_IN
        width_in = height_in * aspect
    if width_in < MIN_SLIDE_IN:
        width_in = MIN_SLIDE_IN
        height_in = width_in / aspect
    if height_in < MIN_SLIDE_IN:
        height_in = MIN_SLIDE_IN
        width_in = height_in * aspect
    # PowerPoint accepts 1–56 inches; extremely elongated pages stay inside.
    width_in = min(56.0, width_in)
    height_in = min(56.0, height_in)
    return Inches(width_in), Inches(height_in)


def _set_east_asian_font(run, font_name: str) -> None:
    """Make sure Korean/CJK text uses the intended typeface in PowerPoint."""

    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        element = rpr.find(qn(tag))
        if element is None:
            element = rpr.makeelement(qn(tag), {})
            rpr.append(element)
        element.set("typeface", font_name)


def _alignment_for_block(block: ExtractedTextBlock) -> PP_ALIGN:
    if len(block.lines) < 2:
        return PP_ALIGN.LEFT
    block_left, block_right = block.bbox[0], block.bbox[2]
    centers = [(line.bbox[0] + line.bbox[2]) / 2 for line in block.lines]
    widths = [line.bbox[2] - line.bbox[0] for line in block.lines]
    block_center = (block_left + block_right) / 2
    centered = all(abs(c - block_center) < 2.0 for c in centers) and max(widths) < (
        block_right - block_left
    ) - 4
    if centered:
        return PP_ALIGN.CENTER
    if all(abs(line.bbox[2] - block_right) < 2.0 for line in block.lines):
        return PP_ALIGN.RIGHT
    return PP_ALIGN.LEFT


class PPTXBuilder:
    """Incrementally adds slides to a presentation."""

    def __init__(self, options: ConversionOptions, page_width: float, page_height: float):
        self.options = options
        self.presentation = Presentation()
        slide_width, slide_height = compute_slide_size(page_width, page_height, options.ratio)
        self.presentation.slide_width = slide_width
        self.presentation.slide_height = slide_height
        self.summary = BuildSummary()

    def _mapping_for(self, page_width: float, page_height: float) -> SlideMapping:
        """Return a mapping that letterboxes one page inside the fixed slide."""

        return SlideMapping(
            self.presentation.slide_width,
            self.presentation.slide_height,
            page_width,
            page_height,
        )

    def _add_warning(self, message: str) -> None:
        if message not in self.summary.warnings:
            self.summary.warnings.append(message)

    def _new_slide(self):
        return self.presentation.slides.add_slide(self.presentation.slide_layouts[6])

    def add_original_page(self, page: "fitz.Page") -> None:
        result = render_page(page, self.options.dpi)
        if result.downscaled:
            self._add_warning(
                f"페이지 {page.number + 1}: 메모리 보호를 위해 렌더링 해상도를 "
                f"{result.dpi} DPI로 낮췄습니다."
            )
        slide = self._new_slide()
        mapping = self._mapping_for(page.rect.width, page.rect.height)
        left, top, width, height = mapping.rect((0, 0, page.rect.width, page.rect.height))
        try:
            slide.shapes.add_picture(io.BytesIO(result.data), left, top, width, height)
        except Exception as exc:
            from app.utils.error_handler import AppError

            raise AppError(
                f"페이지 {page.number + 1}의 이미지를 슬라이드에 삽입하지 못했습니다.",
                detail=str(exc),
            ) from exc

    def add_editable_page(self, doc: "fitz.Document", page_index: int) -> None:
        extracted = extract_page(doc, page_index)
        for warning in extracted.warnings:
            self._add_warning(f"페이지 {page_index + 1}: {warning}")

        page = doc.load_page(page_index)
        slide = self._new_slide()
        mapping = self._mapping_for(extracted.width, extracted.height)

        # If nothing could be rebuilt, keep the page's appearance as an image
        # rather than producing a silently blank slide.
        if not extracted.text_blocks and not extracted.images and not extracted.drawings:
            try:
                has_content = bool(page.get_drawings()) or bool(page.get_text("text").strip())
            except Exception:
                has_content = False
            if has_content:
                self._add_page_image(slide, page, mapping)
                self._add_warning(
                    f"페이지 {page_index + 1}: 편집 가능한 객체를 만들지 못해 이미지로 대체했습니다."
                )
            return

        # Z-order: vector shapes, then images, then text.
        for drawing in extracted.drawings:
            if self._add_drawing(slide, drawing, mapping):
                self.summary.shapes += 1
        for image in extracted.images:
            if self._add_image(slide, image, mapping):
                self.summary.images += 1
        for block in extracted.text_blocks:
            if self._add_text(slide, block, mapping):
                self.summary.text_shapes += 1

    def _add_page_image(self, slide, page: "fitz.Page", mapping: SlideMapping) -> None:
        result = render_page(page, self.options.dpi)
        left, top, width, height = mapping.rect(
            (0, 0, page.rect.width, page.rect.height)
        )
        try:
            slide.shapes.add_picture(io.BytesIO(result.data), left, top, width, height)
        except Exception as exc:
            from app.utils.error_handler import AppError

            raise AppError(
                f"페이지 {page.number + 1}의 이미지를 삽입하지 못했습니다.",
                detail=str(exc),
            ) from exc

    def _apply_drawing_style(self, shape, drawing: ExtractedDrawing) -> None:
        if drawing.fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(*drawing.fill)
        else:
            shape.fill.background()
        if drawing.stroke:
            shape.line.color.rgb = RGBColor(*drawing.stroke)
            shape.line.width = Pt(max(0.25, drawing.stroke_width or 0.75))
        else:
            shape.line.fill.background()

    def _add_drawing(
        self, slide, drawing: ExtractedDrawing, mapping: SlideMapping
    ) -> bool:
        left, top, width, height = mapping.rect(drawing.bbox)

        if drawing.is_rect:
            if width < 1 or height < 1:
                return False
            try:
                shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
                self._apply_drawing_style(shape, drawing)
                return True
            except Exception:
                return False

        points = list(drawing.points)
        if drawing.closed and points and points[0] != points[-1]:
            points.append(points[0])
        if len(points) < 2:
            return False

        # A nearly straight 2-point path is best represented by a connector.
        if len(points) == 2:
            x0, y0 = mapping.x(points[0][0]), mapping.y(points[0][1])
            x1, y1 = mapping.x(points[1][0]), mapping.y(points[1][1])
            if abs(x0 - x1) < 1 and abs(y0 - y1) < 1:
                return False
            try:
                connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x0, y0, x1, y1)
                if drawing.stroke:
                    connector.line.color.rgb = RGBColor(*drawing.stroke)
                    connector.line.width = Pt(max(0.25, drawing.stroke_width or 0.75))
                else:
                    connector.line.fill.background()
                return True
            except Exception:
                return False

        try:
            builder = slide.shapes.build_freeform(
                mapping.x(points[0][0]),
                mapping.y(points[0][1]),
                scale=1.0,
            )
            vertices = [(mapping.x(x), mapping.y(y)) for x, y in points[1:]]
            builder.add_line_segments(vertices, close=drawing.closed)
            shape = builder.convert_to_shape()
            self._apply_drawing_style(shape, drawing)
            return True
        except Exception:
            self._add_warning(
                "복잡한 벡터 도형 일부를 변환하지 못했습니다. 해당 요소는 누락될 수 있습니다."
            )
            return False

    def _add_image(
        self, slide, image: ExtractedImage, mapping: SlideMapping
    ) -> bool:
        left, top, width, height = mapping.rect(image.bbox)
        if width < 2 or height < 2:
            return False
        try:
            slide.shapes.add_picture(io.BytesIO(image.data), left, top, width, height)
            return True
        except Exception:
            self._add_warning("일부 이미지를 변환하지 못했습니다.")
            return False

    def _add_text(
        self, slide, block: ExtractedTextBlock, mapping: SlideMapping
    ) -> bool:
        left, top, width, height = mapping.rect(block.bbox)
        if width < 2 or height < 2:
            return False
        alignment = _alignment_for_block(block)
        try:
            textbox = slide.shapes.add_textbox(left, top, width, height)
            tf = textbox.text_frame
            tf.word_wrap = True
            tf.auto_size = MSO_AUTO_SIZE.NONE
            tf.margin_left = 0
            tf.margin_right = 0
            tf.margin_top = 0
            tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.TOP

            for line_index, line in enumerate(block.lines):
                paragraph = tf.paragraphs[0] if line_index == 0 else tf.add_paragraph()
                paragraph.alignment = alignment
                paragraph.space_before = Pt(0)
                paragraph.space_after = Pt(0)
                paragraph.line_spacing = 1.0
                for span in line.spans:
                    run = paragraph.add_run()
                    run.text = span.text
                    font = run.font
                    font.name = span.font_name
                    font.size = Pt(mapping.font_pt(span.font_size))
                    font.bold = span.bold
                    font.italic = span.italic
                    font.color.rgb = RGBColor(*((span.color >> 16) & 0xFF, (span.color >> 8) & 0xFF, span.color & 0xFF))
                    _set_east_asian_font(run, span.font_name)
            return True
        except Exception:
            self._add_warning("일부 텍스트를 변환하지 못했습니다.")
            return False

    def save(self, path: Path) -> BuildSummary:
        try:
            self.presentation.save(str(path))
        except Exception as exc:
            from app.utils.error_handler import AppError

            raise AppError(
                "PPTX 파일을 저장하지 못했습니다. 저장 공간과 폴더 권한을 확인해 주세요.",
                detail=str(exc),
            ) from exc
        return self.summary

    def discard_last_slide(self) -> None:
        """Remove the most recently added slide after a failed page build."""

        slides = self.presentation.slides
        if len(slides) == 0:
            return
        sld_id_list = slides._sldIdLst
        last = sld_id_list[-1]
        relationship_id = last.get(qn("r:id"))
        sld_id_list.remove(last)
        if relationship_id:
            self.presentation.part.drop_rel(relationship_id)
