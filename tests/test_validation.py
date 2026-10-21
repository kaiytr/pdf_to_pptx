"""Validation, page-range parsing and output-path tests."""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

from app.models.conversion_options import ConversionOptions, parse_page_range
from app.services.validation_service import (
    validate_output_path,
    validate_pptx,
    validate_source_pdf,
)
from app.utils.error_handler import AppError, InvalidPDFError, OutputPathError
from app.utils.file_utils import human_file_size, unique_output_path


# ---------------------------------------------------------------- page ranges
@pytest.mark.parametrize(
    ("spec", "expected"),
    [
        ("", [0, 1, 2, 3, 4]),
        ("   ", [0, 1, 2, 3, 4]),
        ("1", [0]),
        ("2-4", [1, 2, 3]),
        ("1-2,4", [0, 1, 3]),
        ("4,1,2", [3, 0, 1]),
        ("1-2,2-3", [0, 1, 2]),
        (" 1 - 2 , 5 ", [0, 1, 4]),
    ],
)
def test_parse_page_range_valid(spec: str, expected):
    assert parse_page_range(spec, 5) == expected


@pytest.mark.parametrize(
    "spec",
    ["0", "6", "5-7", "3-2", "1--2", "abc", "1-2-3", "-1"],
)
def test_parse_page_range_invalid(spec: str):
    with pytest.raises(ValueError):
        parse_page_range(spec, 5)


def test_parse_page_range_empty_document():
    with pytest.raises(ValueError):
        parse_page_range("", 0)


def test_conversion_options_normalizes_extension(tmp_path: Path):
    options = ConversionOptions(output_path=tmp_path / "result")
    assert options.output_path.suffix == ".pptx"


# ------------------------------------------------------------------- outputs
def test_validate_output_path_accepts_writable_folder(tmp_path: Path):
    path = validate_output_path(tmp_path / "out.pptx")
    assert path == tmp_path / "out.pptx"


def test_validate_output_path_rejects_wrong_extension(tmp_path: Path):
    with pytest.raises(OutputPathError):
        validate_output_path(tmp_path / "out.pdf")


def test_validate_output_path_rejects_missing_folder(tmp_path: Path):
    with pytest.raises(OutputPathError):
        validate_output_path(tmp_path / "missing" / "out.pptx")


def test_validate_output_path_rejects_directory(tmp_path: Path):
    directory = tmp_path / "folder.pptx"
    directory.mkdir()
    with pytest.raises(OutputPathError):
        validate_output_path(directory)


def _make_pptx(path: Path, slide_count: int = 1) -> Path:
    presentation = Presentation()
    for _ in range(slide_count):
        presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(str(path))
    return path


def test_validate_pptx_accepts_generated_file(tmp_path: Path):
    path = _make_pptx(tmp_path / "valid.pptx", 2)
    validate_pptx(path, expected_slides=2)


def test_validate_pptx_rejects_missing_file(tmp_path: Path):
    with pytest.raises(AppError):
        validate_pptx(tmp_path / "missing.pptx", expected_slides=1)


def test_validate_pptx_rejects_non_zip(tmp_path: Path):
    path = tmp_path / "fake.pptx"
    path.write_bytes(b"this is not a zip file")
    with pytest.raises(AppError):
        validate_pptx(path, expected_slides=1)


def test_validate_pptx_rejects_wrong_slide_count(tmp_path: Path):
    path = _make_pptx(tmp_path / "valid.pptx", 1)
    with pytest.raises(AppError):
        validate_pptx(path, expected_slides=3)


def test_validate_pptx_rejects_missing_parts(tmp_path: Path):
    path = tmp_path / "partial.pptx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
    with pytest.raises(AppError):
        validate_pptx(path, expected_slides=1)


def test_validate_source_pdf(sample_pdf: Path):
    info = validate_source_pdf(sample_pdf)
    assert info.page_count == 3


def test_validate_source_pdf_rejects_corrupt(corrupt_pdf: Path):
    with pytest.raises((InvalidPDFError, AppError)):
        validate_source_pdf(corrupt_pdf)


# --------------------------------------------------------------------- utils
def test_human_file_size():
    assert human_file_size(0) == "0 B"
    assert human_file_size(1024) == "1.0 KB"
    assert human_file_size(5 * 1024 * 1024) == "5.0 MB"


def test_unique_output_path(tmp_path: Path):
    first = unique_output_path(tmp_path, "deck")
    assert first.name == "deck.pptx"
    first.write_bytes(b"x")
    second = unique_output_path(tmp_path, "deck")
    assert second.name == "deck (2).pptx"
    second.write_bytes(b"x")
    assert unique_output_path(tmp_path, "deck").name == "deck (3).pptx"
